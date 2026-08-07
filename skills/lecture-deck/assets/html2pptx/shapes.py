"""python-pptx 저수준 헬퍼.

python-pptx 는 채우기/선에 알파를 직접 노출하지 않는다. theme.css 의 카드가
rgba(244,244,238,.03) 처럼 반투명이라 배경 글로우가 비쳐야 하므로,
solidFill 안에 <a:alpha> 를 직접 꽂아 넣는다.
"""

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

from theme import EMU_PER_PX, px

# --------------------------------------------------------------------------
# 채우기 / 선
# --------------------------------------------------------------------------


def _apply_alpha(srgb_el, alpha: float):
    """<a:srgbClr> 아래에 <a:alpha val="..."/> 를 꽂는다."""
    from lxml import etree

    if srgb_el is None:
        return
    el = etree.SubElement(srgb_el, qn("a:alpha"))
    el.set("val", str(int(round(alpha * 100000))))


def _fill_parent(shape):
    """도형이면 spPr, 표 셀이면 tcPr — solidFill 이 들어가는 부모 엘리먼트."""
    el = getattr(shape, "_element", None)
    if el is not None and hasattr(el, "spPr"):
        return el.spPr
    tc = getattr(shape, "_tc", None)
    if tc is not None:
        return tc.get_or_add_tcPr()
    return None


def fill_rgba(shape, hex_color: str, alpha: float = 1.0):
    """도형(또는 표 셀) 면을 hex_color + alpha 로 채운다."""
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(hex_color)
    if alpha < 1.0:
        parent = _fill_parent(shape)
        if parent is not None:
            _apply_alpha(parent.find(qn("a:solidFill")).find(qn("a:srgbClr")), alpha)


def style_table_dark(table, line_hex: str, line_alpha: float, width_px: float = 1.0):
    """PowerPoint 기본 표 스타일(하늘색 밴딩)을 걷어내고 헤어라인 테두리를 준다."""
    tbl = table._tbl
    tbl_pr = tbl.find(qn("a:tblPr"))
    if tbl_pr is not None:
        for attr, val in (("firstRow", "0"), ("bandRow", "0"), ("bandCol", "0"),
                          ("firstCol", "0")):
            tbl_pr.set(attr, val)
        style_id = tbl_pr.find(qn("a:tableStyleId"))
        if style_id is not None:
            tbl_pr.remove(style_id)

    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            # 스키마 순서: lnL, lnR, lnT, lnB 가 tcPr 의 맨 앞에 와야 한다
            for i, tag in enumerate(("a:lnB", "a:lnT", "a:lnR", "a:lnL")):
                ln = tc_pr.makeelement(qn(tag), {"w": str(px(width_px)), "cap": "flat"})
                sf = ln.makeelement(qn("a:solidFill"), {})
                clr = sf.makeelement(qn("a:srgbClr"), {"val": line_hex})
                _apply_alpha(clr, line_alpha)
                sf.append(clr)
                ln.append(sf)
                tc_pr.insert(0, ln)


def no_fill(shape):
    shape.fill.background()


def line_rgba(shape, hex_color: str, alpha: float = 1.0, width_px: float = 1.0):
    shape.line.color.rgb = RGBColor.from_string(hex_color)
    shape.line.width = px(width_px)
    if alpha < 1.0:
        ln = shape.line._get_or_add_ln()
        _apply_alpha(ln.find(qn("a:solidFill")).find(qn("a:srgbClr")), alpha)


def no_line(shape):
    shape.line.fill.background()


def soft_shadow(shape, hex_color="503C28", alpha=0.13, blur_px=16, dy_px=6):
    """카드에 낮고 부드러운 그림자를 준다.

    python-pptx 는 그림자를 노출하지 않으므로 <a:effectLst><a:outerShdw> 를 직접
    넣는다. 밝은 테마에서 흰 카드를 배경에서 떼어내는 유일한 수단이다.
    """
    from lxml import etree

    sp_pr = shape._element.spPr
    for old in sp_pr.findall(qn("a:effectLst")):
        sp_pr.remove(old)
    eff = sp_pr.makeelement(qn("a:effectLst"), {})
    shdw = etree.SubElement(eff, qn("a:outerShdw"), {
        "blurRad": str(px(blur_px)),
        "dist": str(px(dy_px)),
        "dir": "5400000",          # 90도 = 아래로
        "rotWithShape": "0",
    })
    clr = etree.SubElement(shdw, qn("a:srgbClr"), {"val": hex_color})
    _apply_alpha(clr, alpha)
    # effectLst 는 스키마상 채우기·선 뒤에 온다
    sp_pr.append(eff)
    shape.shadow.inherit = False


def set_theme_hyperlink_color(prs, hex_color: str):
    """테마의 하이퍼링크 색을 바꾼다.

    PowerPoint 는 하이퍼링크 런의 글자색을 rPr 의 solidFill 이 아니라 테마의
    a:hlink 로 칠한다. 런 색만 지정하면 계속 파란 링크로 나오므로 테마를 고친다.
    """
    import re

    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    for master in prs.slide_masters:
        part = master.part.part_related_by(RT.THEME)
        xml = part.blob.decode("utf-8")
        for tag in ("hlink", "folHlink"):
            xml = re.sub(
                rf"<a:{tag}>.*?</a:{tag}>",
                f'<a:{tag}><a:srgbClr val="{hex_color}"/></a:{tag}>',
                xml,
                count=1,
            )
        part._blob = xml.encode("utf-8")


# --------------------------------------------------------------------------
# 도형
# --------------------------------------------------------------------------


def card(slide, x, y, w, h, radius=4, fill=None, fill_alpha=1.0,
         border=None, border_alpha=1.0, border_px=1.0):
    """헤어라인 카드. x/y/w/h 는 CSS px."""
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, px(x), px(y), px(w), px(h)
    )
    # adjustment 는 짧은 변 대비 비율
    shp.adjustments[0] = min(0.5, radius / max(1.0, min(w, h)))
    if fill:
        fill_rgba(shp, fill, fill_alpha)
    else:
        no_fill(shp)
    if border:
        line_rgba(shp, border, border_alpha, border_px)
    else:
        no_line(shp)
    shp.shadow.inherit = False
    return shp


def bar(slide, x, y, w, h, color, alpha=1.0):
    """테두리 강조선·코너 브래킷용 실선 사각형."""
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(x), px(y), px(w), px(h))
    fill_rgba(shp, color, alpha)
    no_line(shp)
    shp.shadow.inherit = False
    return shp


def corner_brackets(slide, x, y, w, h, color, arm=16, thick=2):
    """case-card 의 시그니처: 좌상 ┌ , 우하 ┘ 그린 브래킷."""
    bar(slide, x, y, arm, thick, color)
    bar(slide, x, y, thick, arm, color)
    bar(slide, x + w - arm, y + h - thick, arm, thick, color)
    bar(slide, x + w - thick, y + h - arm, thick, arm, color)


def pill(slide, x, y, w, h, fill=None, fill_alpha=1.0, border=None, border_alpha=1.0):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(x), px(y), px(w), px(h))
    shp.adjustments[0] = 0.5
    if fill:
        fill_rgba(shp, fill, fill_alpha)
    else:
        no_fill(shp)
    if border:
        line_rgba(shp, border, border_alpha, 1.0)
    else:
        no_line(shp)
    shp.shadow.inherit = False
    return shp


def hline(slide, x, y, w, color, alpha=1.0, thick=1):
    return bar(slide, x, y, w, thick, color, alpha)


# --------------------------------------------------------------------------
# 텍스트
# --------------------------------------------------------------------------

ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


def textbox(slide, x, y, w, h, anchor="top", wrap=True):
    """여백 0 인 텍스트 상자. x/y/w/h 는 CSS px."""
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = ANCHOR[anchor]
    # 기본 단락 하나가 딸려오므로 재사용한다
    tf.paragraphs[0].text = ""
    return tb


def para(tf, first=False, align="left", line_spacing=None, space_before=0, space_after=0):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = ALIGN[align]
    if line_spacing:
        p.line_spacing = line_spacing
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    return p


def run(p, text, font, size_px, color, bold=False, italic=False, link=None,
        underline=False, strike=False, spacing=None):
    """단락에 런 하나를 붙인다. size_px 는 CSS px, spacing 은 letter-spacing(px)."""
    r = p.add_run()
    r.text = text
    f = r.font
    f.name = font
    f.size = Pt(size_px * 0.75)
    f.bold = bold
    f.italic = italic
    f.underline = underline
    f.color.rgb = RGBColor.from_string(color)
    if link:
        r.hyperlink.address = link
        # 하이퍼링크는 테마색으로 덮이므로 색을 다시 강제한다
        f.color.rgb = RGBColor.from_string(color)
    rPr = r._r.get_or_add_rPr()
    if strike:
        rPr.set("strike", "sngStrike")
    if spacing:
        rPr.set("spc", str(int(round(spacing * 0.75 * 100))))
    _set_latin_and_ea(r, font)
    return r


def line_break(p):
    """단락 안 줄바꿈 <a:br/>. endParaRPr 앞에 끼워 넣어야 한다."""
    br = p._p.makeelement(qn("a:br"), {})
    end = p._p.find(qn("a:endParaRPr"))
    if end is not None:
        end.addprevious(br)
    else:
        p._p.append(br)
    return br


def _set_latin_and_ea(r, font_name: str):
    """한글이 라틴 폰트로 떨어지지 않도록 ea/cs 까지 같은 폰트로 고정.

    a:ea / a:cs 는 python-pptx 가 모르는 자식이라 그냥 append 하면 rPr 맨 뒤,
    즉 a:hlinkClick 뒤에 붙는다. 스키마 순서(latin → ea → cs → hlinkClick)를
    어기면 PowerPoint 가 그 런의 rPr 을 통째로 무시하고 하이퍼링크를 파란 밑줄
    기본값으로 그려버린다. 반드시 a:latin 바로 뒤에 끼워 넣어야 한다.
    """
    rPr = r._r.get_or_add_rPr()
    prev = rPr.get_or_add_latin()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            prev.addnext(el)
        el.set("typeface", font_name)
        prev = el


def spacing_px(v: float) -> float:
    """CSS px 간격 → Pt (space_before/after 용)."""
    return v * 0.75


def emu_px(v: float) -> int:
    return int(round(v * EMU_PER_PX))
