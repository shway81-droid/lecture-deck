"""블록 IR → PPTX 도형.

블록마다 measure(높이 계산)와 draw(그리기) 한 쌍을 둔다. python-pptx 에는
자동 레이아웃이 없으므로 build.py 가 measure 로 총 높이를 먼저 구해
수직 중앙 정렬한 뒤 draw 를 부른다.

좌표 단위는 전부 CSS px (1280x720 기준)이고, S 는 슬라이드가 넘칠 때
reveal 의 shrink-to-fit 처럼 전체를 줄이는 축소 계수다.
"""

from pathlib import Path

from PIL import Image
from pptx.util import Pt

import theme as T
from parse import ACCENT, BHOT, BR, CITATION, CODE, OP, PLAIN, STRONG
from shapes import (bar, card, corner_brackets, fill_rgba, line_break, para,
                    pill, px, run, soft_shadow, style_table_dark, textbox)


def ls(base_px, lh):
    """정확한 줄간격(Pt).

    PowerPoint 에서 line_spacing 을 배수로 주면 '폰트 기본 줄높이'에 곱해진다.
    Noto Sans KR 은 기본 줄높이가 1.4em 쯤이라 CSS line-height 와 어긋나고,
    측정값과 실제 조판이 벌어져 블록끼리 겹친다. 그래서 항상 절대값으로 준다.
    """
    return Pt(base_px * lh * 0.75)


def _has_hangul(s: str) -> bool:
    return any(0xAC00 <= ord(c) <= 0xD7A3 or 0x1100 <= ord(c) <= 0x11FF for c in s)


# --------------------------------------------------------------------------
# 런 스타일 해석
# --------------------------------------------------------------------------


def run_style(style, base_px, S, ctx_font, ctx_color, text="", ctx_bold=False,
              strong_color=None):
    """(font, size_px, color, bold, italic) 를 돌려준다.

    bold 는 최종값이다. .accent 는 CSS 에서 font-weight:400 이라 h1 안에 있어도
    굵어지면 안 되므로 문맥 굵기를 물려받지 않는다.
    strong_color 는 .hero-card .big strong 처럼 strong 색만 바뀌는 자리에 쓴다.
    """
    if style == STRONG:
        return ctx_font, base_px, strong_color or T.FG_STRONG, True, False
    if style == ACCENT:
        if T.STYLE["accent_run"] == "bold":
            # 강조를 볼드로 처리하는 테마 (Warm Paper). 이탤릭·세리프를 쓰지 않는다.
            return ctx_font, base_px, T.ACCENT, True, False
        # Instrument Serif 에는 한글 글리프가 없어 브라우저는 Pretendard 이탤릭으로
        # 떨어진다. 그 동작을 그대로 흉내낸다 — 한글이면 sans, 라틴이면 serif.
        font = T.SANS if _has_hangul(text) else T.SERIF
        return font, base_px, T.ACCENT, False, True
    if style == CODE:
        return T.MONO, base_px * 0.85, T.FG_STRONG, ctx_bold, False
    if style == CITATION:
        return T.SANS, T.FS["citation"] * S, T.ACCENT, False, False
    if style == OP:
        return ctx_font, base_px, T.MUTED, ctx_bold, False
    if style == BHOT:
        return ctx_font, base_px, T.ACCENT, True, False
    return ctx_font, base_px, ctx_color, ctx_bold, False


def rich_lines(runs, base_px, S, box_w, ctx_font=None, ctx_color=None):
    """서식이 섞인 런 목록이 box_w 안에서 차지하는 줄 수."""
    ctx_font = ctx_font or T.SANS
    ctx_color = ctx_color or T.TEXT
    lines, used = 1, 0.0
    for r in runs:
        if r["style"] == BR:
            lines += 1
            used = 0.0
            continue
        _, size, _, _, _ = run_style(r["style"], base_px, S, ctx_font, ctx_color,
                                     r["text"])
        for tok in T._tokenize(r["text"]):
            w = T.text_width_px(tok, size)
            if used > 0 and used + w > box_w:
                lines += 1
                used = 0.0 if tok.strip() == "" else w
            else:
                used += w
    return lines


def rich_h(runs, base_px, S, box_w, lh=1.65, ctx_font=None):
    return rich_lines(runs, base_px, S, box_w, ctx_font) * base_px * lh


def rich_text(slide, x, y, w, h, runs, base_px, S, *, font=None, color=None,
              lh=1.65, align="left", anchor="top", bold=False, spacing=None,
              strike=False, strong_color=None, wrap=True):
    """런 목록을 텍스트 상자 하나에 그린다."""
    font = font or T.SANS
    color = color or T.TEXT
    tb = textbox(slide, x, y, w, h, anchor=anchor, wrap=wrap)
    p = para(tb.text_frame, first=True, align=align, line_spacing=ls(base_px, lh))
    for r in runs:
        if r["style"] == BR:
            line_break(p)
            continue
        f, size, col, b, it = run_style(r["style"], base_px, S, font, color,
                                        r["text"], ctx_bold=bold,
                                        strong_color=strong_color)
        run(p, r["text"], f, size, col, bold=b, italic=it,
            link=r.get("link"), spacing=spacing, strike=strike)
    return tb


def plain_text(slide, x, y, w, h, text, size_px, color, *, font=None, lh=1.65,
               align="left", anchor="top", bold=False, spacing=None, strike=False,
               italic=False, link=None):
    tb = textbox(slide, x, y, w, h, anchor=anchor)
    p = para(tb.text_frame, first=True, align=align, line_spacing=ls(size_px, lh))
    run(p, text, font or T.SANS, size_px, color, bold=bold, italic=italic,
        spacing=spacing, strike=strike, link=link)
    return tb


# --------------------------------------------------------------------------
# 테마별 카드 처리
# --------------------------------------------------------------------------


def card_style():
    """표준 카드의 면·테두리 인자. 테마가 헤어라인이냐 그림자냐를 정한다."""
    if T.STYLE["card"] == "shadow":
        return dict(fill=T.SURFACE, fill_alpha=T.A_SURFACE_1,
                    border=T.RULE_C, border_alpha=0.0)
    return dict(fill=T.SURFACE, fill_alpha=T.A_SURFACE_1,
                border=T.RULE_C, border_alpha=T.A_RULE)


def tcard(slide, x, y, w, h, radius=None, **kw):
    """테마 카드 하나. shadow 테마면 부드러운 그림자를 얹는다."""
    kw = {**card_style(), **kw}
    shp = card(slide, x, y, w, h,
               radius=T.RADIUS_SM if radius is None else radius, **kw)
    if T.STYLE["card"] == "shadow":
        soft_shadow(shp)
    return shp


# --------------------------------------------------------------------------
# 블록별 바깥 여백 (margin-top, margin-bottom) — CSS 를 px 로 옮긴 값
# --------------------------------------------------------------------------

MARGIN = {
    "h1": (0, 17.6),
    "h2": (0, 25.4),
    "h2.compact": (0, 15.0),
    "h2.small": (0, 17.0),
    "eyebrow": (0, 9.5),
    "p": (20, 20),
    "pull-quote": (26.4, 18.5),
    "kpi-line": (12.8, 12.8),
    "formula": (14, 2),
    "bullets": (8.8, 8.8),
    "tl": (14, 14),
    "card-row": (10, 0),
    "dual": (10, 0),
    "thesis-flow": (22, 22),
    "cap-grid": (16, 16),
    "ladder": (14, 14),
    "check-list": (12, 12),
    "stat-row": (18, 18),
    "case-card": (10, 10),
    "brief": (10, 10),
    "code": (20, 20),
    "source-list": (11, 0),
    "title-block": (0, 0),
    "card-grid": (16, 16),
    "table": (18, 18),
    "tier-grid": (18, 18),
    "chip-row": (14, 14),
    "qr-grid": (24, 16),
    "image": (16, 16),
    "section-num": (0, 4.4),
    "byline": (-4, 6),
    "hero-sub": (0, 0),
    "beat-sub": (0, 0),
}


def margin_key(b) -> str:
    k = b["kind"]
    if k == "h2":
        return {"compact": "h2.compact", "small": "h2.small",
                "beat": "h2"}.get(b.get("variant"), "h2")
    if k == "p":
        for c in ("eyebrow", "section-num", "byline", "pull-quote", "kpi-line",
                  "formula", "hero-sub", "beat-sub"):
            if c in b["classes"]:
                return c
        return "p"
    return k


def margins(b, S):
    mt, mb = MARGIN.get(margin_key(b), (12, 12))
    return mt * S, mb * S


# --------------------------------------------------------------------------
# p — eyebrow / muted / small / center / pull-quote / kpi-line / formula / …
# --------------------------------------------------------------------------


def _p_spec(b, S):
    """p 계열의 (font, size_px, color, lh, align, letter_spacing, bold) 결정."""
    c = b["classes"]
    if "eyebrow" in c or "section-num" in c:
        size = T.FS["section_num"] * S if "section-num" in c else T.FS["eyebrow"] * S
        return (T.MONO, size, T.ACCENT, 1.4, T.STYLE["eyebrow_align"],
                0.2 * size, "section-num" in c)
    if "byline" in c:
        return T.MONO, T.FS["byline"] * S, T.MUTED, 1.4, "left", None, False
    if "hero-sub" in c:
        return T.SANS, T.FS["hero_sub"] * S, T.TEXT, T.LH["hero_sub"], "left", None, False
    if "beat-sub" in c:
        return T.SANS, T.FS["beat_sub"] * S, T.FG2, T.LH["beat_sub"], "left", None, False
    if "pull-quote" in c:
        return T.SANS, T.FS["pull_quote"] * S, T.FG_STRONG, T.LH["pull_quote"], "center", None, True
    if "kpi-line" in c:
        return T.SANS, T.FS["kpi"] * S, T.FG_STRONG, T.LH["kpi"], "left", None, True
    if "formula" in c:
        return T.MONO, T.FS["formula"] * S, T.MUTED, 1.4, "center", None, False
    if "big" in c:
        return T.SANS, T.FS["hero_big"] * S, T.TEXT, T.LH["hero_big"], "left", None, False
    if "subtitle" in c:
        return T.SANS, T.FS["subtitle"] * S, T.MUTED, T.LH["subtitle"], "left", None, False

    size = T.FS["small"] * S if "small" in c else T.FS["body"] * S
    color = T.MUTED if "muted" in c else T.TEXT
    align = "center" if "center" in c else "left"
    return T.SANS, size, color, T.LH["body"], align, None, False


def m_p(b, w, S):
    font, size, _, lh, _, _, _ = _p_spec(b, S)
    runs = b["runs"]
    if "pull-quote" in b["classes"]:
        w = min(w, 900 * S)
        runs = _quoted(runs)
    return rich_lines(runs, size, S, w, font) * size * lh


def _quoted(runs):
    """.pull-quote::before/after — 그린 세리프 따옴표."""
    return (
        [{"text": "“", "style": ACCENT, "link": None}]
        + runs
        + [{"text": "”", "style": ACCENT, "link": None}]
    )


def d_p(slide, b, x, y, w, S):
    font, size, color, lh, align, spc, bold = _p_spec(b, S)
    runs = b["runs"]
    if "eyebrow" in b["classes"] and T.STYLE["eyebrow_prefix"]:
        runs = [{"text": T.STYLE["eyebrow_prefix"], "style": PLAIN, "link": None}] + runs
    if "pull-quote" in b["classes"]:
        runs = _quoted(runs)
        inner = min(w, 900 * S)
        x += (w - inner) / 2
        w = inner
    h = m_p(b, w, S)
    rich_text(slide, x, y, w, h, runs, size, S, font=font, color=color,
              lh=lh, align=align, bold=bold, spacing=spc)


# --------------------------------------------------------------------------
# h1 / h2
# --------------------------------------------------------------------------


def _h_size(b, S):
    v = b.get("variant")
    if b["kind"] == "h1":
        if v == "divider":
            return T.FS["h1_divider"] * S, T.LH["h1"]
        return T.FS["h1"] * S, T.LH["h1"]
    if v == "compact":
        return T.FS["h2_compact"] * S, T.LH["h2"]
    if v == "small":
        return T.FS["h2_small"] * S, T.LH["h2"]
    if v == "beat":
        return T.FS["beat_h2"] * S, T.LH["beat_h2"]
    return T.FS["h2"] * S, T.LH["h2"]


def m_h(b, w, S):
    size, lh = _h_size(b, S)
    return rich_lines(b["runs"], size, S, w, T.SANS) * size * lh


def d_h(slide, b, x, y, w, S):
    size, lh = _h_size(b, S)
    color = T.MUTED if b.get("variant") == "small" else T.FG_STRONG
    bold = b.get("variant") != "small"
    # h1(표지)은 항상 왼쪽. h2 는 테마가 정한다.
    align = "left" if b["kind"] == "h1" else T.STYLE["h2_align"]
    rich_text(slide, x, y, w, m_h(b, w, S), b["runs"], size, S,
              font=T.SANS if not bold else T.SANS_BLACK if b["kind"] == "h1" else T.SANS,
              color=color, lh=lh, align=align, bold=bold)


# --------------------------------------------------------------------------
# bullets (▸ 마커)
# --------------------------------------------------------------------------


def m_bullets(b, w, S):
    size = T.FS["body"] * S
    ind = T.BULLET_INDENT * S
    gap = T.BULLET_GAP * S
    total = 0.0
    for i, item in enumerate(b["items"]):
        total += rich_lines(item, size, S, w - ind) * size * T.LH["body"]
        if i:
            total += gap
    return total


def _bullet_marker(b, i):
    """.bullets 는 theme.css 의 그린 ▸, 그 밖의 목록은 브라우저 기본 마커."""
    if b.get("ordered"):
        return f"{i + 1}.", T.ACCENT
    if b.get("styled", True):
        return "▸", T.ACCENT
    return "•", T.FG2


def d_bullets(slide, b, x, y, w, S):
    size = T.FS["body"] * S
    ind = T.BULLET_INDENT * S
    gap = T.BULLET_GAP * S
    cy = y
    for i, item in enumerate(b["items"]):
        h = rich_lines(item, size, S, w - ind) * size * T.LH["body"]
        mark, mcolor = _bullet_marker(b, i)
        plain_text(slide, x, cy, ind, size * T.LH["body"], mark, size,
                   mcolor, bold=True, lh=T.LH["body"])
        rich_text(slide, x + ind, cy, w - ind, h, item, size, S,
                  color=T.FG2, lh=T.LH["body"])
        cy += h + (gap if i < len(b["items"]) - 1 else 0)
    return cy - y


# --------------------------------------------------------------------------
# tl — 연표
# --------------------------------------------------------------------------


def _tl_yr_w(b, S):
    """연도 칸 너비.

    CSS 는 58px 고정이라 '2025.04' 처럼 긴 값이 본문 위로 삐져나가 겹친다
    (원본 브라우저에서도 겹친다). 칸을 실제 글자 폭까지 넓혀 그 결함을 없앤다.
    """
    yr_size = T.FS["tl_yr"] * S
    widest = max(
        (T.text_width_px(_plain_runs(r["yr"]), yr_size) for r in b["rows"]),
        default=0.0,
    )
    return max(T.TL_YR_W * S, widest + 4 * S)


def _plain_runs(runs) -> str:
    return "".join("" if r["style"] == BR else r["text"] for r in runs)


def _tl_row_h(row, w, S, yr_w):
    size = T.FS["tl_ev"] * S
    ev_w = w - yr_w - T.TL_COL_GAP * S
    return max(
        rich_lines(row["ev"], size, S, ev_w) * size * T.LH["tl_ev"],
        size * T.LH["tl_ev"],
    ) + T.TL_ROW_PAD_Y * 2 * S


def m_tl(b, w, S):
    yr_w = _tl_yr_w(b, S)
    total = sum(_tl_row_h(r, w, S, yr_w) for r in b["rows"])
    total += T.TL_GAP * S * max(0, len(b["rows"]) - 1)
    total += 1 * len(b["rows"])          # 행마다 하단 헤어라인
    return total


def d_tl(slide, b, x, y, w, S):
    size = T.FS["tl_ev"] * S
    yr_size = T.FS["tl_yr"] * S
    yr_w = _tl_yr_w(b, S)
    ev_x = x + yr_w + T.TL_COL_GAP * S
    ev_w = w - yr_w - T.TL_COL_GAP * S
    cy = y
    for row in b["rows"]:
        h = _tl_row_h(row, w, S, yr_w)
        ty = cy + T.TL_ROW_PAD_Y * S
        rich_text(slide, x + 4 * S, ty, yr_w, yr_size * 1.4, row["yr"],
                  yr_size, S, font=T.MONO, color=T.ACCENT, lh=T.LH["tl_ev"],
                  bold=True, wrap=False)
        rich_text(slide, ev_x, ty, ev_w, h, row["ev"], size, S,
                  color=T.FG_STRONG if row["now"] else T.FG2, lh=T.LH["tl_ev"])
        if row["now"]:
            bar(slide, x, cy + h, w, 1, T.ACCENT)
        else:
            bar(slide, x, cy + h, w, 1, T.RULE_C, T.A_RULE)
        cy += h + 1 + T.TL_GAP * S
    return cy - y


# --------------------------------------------------------------------------
# dual — 좌우 비교 카드
# --------------------------------------------------------------------------


def _dual_plain(col) -> bool:
    """.dd 없이 라벨+불릿만 있는 layouts.md 형태인가."""
    return col["tone"] == "plain"


def _dual_col_h(col, cw, S):
    plain_col = _dual_plain(col)
    pad_x = 0 if plain_col else T.DUAL_PAD_X * S
    pad_y = 0 if plain_col else T.DUAL_PAD_Y * S
    hd = (T.FS["eyebrow"] if plain_col else T.FS["dd_h"]) * S
    li = (T.FS["body"] if plain_col else T.FS["dd_li"]) * S
    lh = T.LH["body"] if plain_col else T.LH["dd_li"]
    marker = (T.BULLET_INDENT if plain_col else 18) * S
    inner = cw - pad_x * 2 - marker
    header_bar = T.STYLE["dual"] == "header-bar" and not plain_col
    if not col["head"]:
        h = 0.0
    elif header_bar:
        h = hd * 1.4 + 22 * S + 12 * S      # 채움 헤더 + 아래 여백
    else:
        h = hd * 1.4 + 10 * S
    for i, item in enumerate(col["items"]):
        h += rich_lines(item, li, S, inner) * li * lh
        if i:
            h += (T.BULLET_GAP if plain_col else 7) * S
    for extra in col.get("extra", []):
        h += rich_lines(extra, li, S, cw - pad_x * 2) * li * lh + 6 * S
    return h + pad_y * 2


def m_dual(b, w, S):
    cw = (w - T.DUAL_GAP * S) / 2
    return max(_dual_col_h(c, cw, S) for c in b["cols"])


def d_dual(slide, b, x, y, w, S):
    cw = (w - T.DUAL_GAP * S) / 2
    h = m_dual(b, w, S)

    for i, col in enumerate(b["cols"]):
        cx = x + i * (cw + T.DUAL_GAP * S)
        plain_col = _dual_plain(col)
        pad_x = 0 if plain_col else T.DUAL_PAD_X * S
        pad_y = 0 if plain_col else T.DUAL_PAD_Y * S
        hd = (T.FS["eyebrow"] if plain_col else T.FS["dd_h"]) * S
        li = (T.FS["body"] if plain_col else T.FS["dd_li"]) * S
        lh = T.LH["body"] if plain_col else T.LH["dd_li"]
        marker = (T.BULLET_INDENT if plain_col else 18) * S

        header_bar = T.STYLE["dual"] == "header-bar" and not plain_col
        head_h = hd * 1.4 + (22 * S if header_bar else 0)

        if not plain_col:
            tcard(slide, cx, y, cw, h, radius=T.RADIUS_SM * S)
            if header_bar:
                # 상단 채움 헤더: 금지·제약 쪽은 잉크, 허용·권장 쪽은 보조색
                hb = T.ACCENT2 if col["tone"] == "do" else T.INK
                bar(slide, cx, y, cw, head_h, hb)
            elif T.STYLE["dual"] == "left-bar":
                # 긍정/부정을 색으로 가른다 (dotpaper). 코랄이 없는 테마는 헤어라인으로.
                tone_c = T.ACCENT2 if col["tone"] == "do" else getattr(T, "ACCENT3", T.RULE_C)
                bar(slide, cx, y, 4 * S, h, tone_c)
            elif col["tone"] == "do":
                bar(slide, cx, y, 2 * S, h, T.ACCENT)
            else:
                bar(slide, cx, y, 2 * S, h, T.RULE_C, T.A_RULE_STRONG)

        tx = cx + pad_x
        tw = cw - pad_x * 2
        cy = y + pad_y

        if col["head"]:
            head_runs = col["head"]
            if plain_col:      # eyebrow ::before "— "
                head_runs = [{"text": "— ", "style": PLAIN, "link": None}] + head_runs
            if header_bar:
                rich_text(slide, cx, y, cw, head_h, head_runs, hd, S,
                          font=T.SANS, color="FFFFFF", lh=1.4, align="center",
                          anchor="middle", bold=True)
                cy = y + head_h + 12 * S
            else:
                rich_text(slide, tx, cy, tw, hd * 1.4, head_runs, hd, S, font=T.MONO,
                          color=T.MUTED if col["tone"] == "dont" else T.ACCENT,
                          lh=1.4, spacing=(0.2 if plain_col else 0.14) * hd,
                          strike=col["tone"] == "dont")
                cy += hd * 1.4 + 10 * S

        li_color = T.MUTED if col["tone"] == "dont" else T.FG2
        mark = "▸" if plain_col else "•"
        gap = (T.BULLET_GAP if plain_col else 7) * S
        for j, item in enumerate(col["items"]):
            ih = rich_lines(item, li, S, tw - marker) * li * lh
            plain_text(slide, tx, cy, marker, li * lh, mark, li,
                       T.ACCENT if plain_col else li_color, lh=lh,
                       bold=plain_col)
            rich_text(slide, tx + marker, cy, tw - marker, ih, item, li, S,
                      color=li_color, lh=lh)
            cy += ih + (gap if j < len(col["items"]) - 1 else 0)

        for extra in col.get("extra", []):
            eh = rich_lines(extra, li, S, tw) * li * lh
            rich_text(slide, tx, cy + 6 * S, tw, eh, extra, li, S,
                      color=li_color, lh=lh)
            cy += eh + 6 * S
    return h


# --------------------------------------------------------------------------
# card-row — hero 카드 + 불릿
# --------------------------------------------------------------------------


def m_card_row(b, w, S):
    cw = (w - T.DUAL_GAP * S) / 2
    big = T.FS["hero_big"] * S
    hero_h = (rich_lines(b["hero"], big, S, cw - T.HERO_PAD * 2 * S) * big
              * T.LH["hero_big"] + T.HERO_PAD * 2 * S)
    bullets_h = m_bullets({"items": b["items"]}, cw, S)
    return max(hero_h, bullets_h)


def d_card_row(slide, b, x, y, w, S):
    cw = (w - T.DUAL_GAP * S) / 2
    big = T.FS["hero_big"] * S
    inner = cw - T.HERO_PAD * 2 * S
    hero_text_h = rich_lines(b["hero"], big, S, inner) * big * T.LH["hero_big"]
    hero_h = hero_text_h + T.HERO_PAD * 2 * S
    bullets_h = m_bullets({"items": b["items"]}, cw, S)
    h = max(hero_h, bullets_h)

    hy = y + (h - hero_h) / 2
    card(slide, x, hy, cw, hero_h, radius=T.RADIUS_LG * S,
         **card_style())
    # .hero-card .big strong { color: accent } — 색만 그린, 굵기와 서체는 그대로
    rich_text(slide, x + T.HERO_PAD * S, hy + T.HERO_PAD * S, inner, hero_text_h,
              b["hero"], big, S, lh=T.LH["hero_big"], strong_color=T.ACCENT)

    by = y + (h - bullets_h) / 2
    d_bullets(slide, {"items": b["items"]}, x + cw + T.DUAL_GAP * S, by, cw, S)
    return h


# --------------------------------------------------------------------------
# thesis-flow
# --------------------------------------------------------------------------


def _flow_h(step, cw, S):
    t = T.FS["fs_t"] * S
    n = T.FS["fs_n"] * S
    inner = cw - T.FLOW_PAD_X * 2 * S
    return (t * 1.4 + 8 * S
            + rich_lines(step["n"], n, S, inner) * n * T.LH["fs_n"]
            + T.FLOW_PAD_Y * 2 * S)


def m_thesis_flow(b, w, S):
    k = len(b["steps"])
    cw = (w - T.FLOW_GAP * S * (k - 1)) / k
    return max(_flow_h(s, cw, S) for s in b["steps"])


def d_thesis_flow(slide, b, x, y, w, S):
    k = len(b["steps"])
    cw = (w - T.FLOW_GAP * S * (k - 1)) / k
    h = m_thesis_flow(b, w, S)
    t = T.FS["fs_t"] * S
    n = T.FS["fs_n"] * S
    for i, step in enumerate(b["steps"]):
        cx = x + i * (cw + T.FLOW_GAP * S)
        if step["now"] and T.STYLE["card"] == "shadow":
            tcard(slide, cx, y, cw, h, radius=T.RADIUS_SM * S,
                  fill=T.ACCENT, fill_alpha=T.A_ACCENT_06,
                  border=T.ACCENT, border_alpha=1.0)
        else:
            tcard(slide, cx, y, cw, h, radius=T.RADIUS_SM * S,
                  border=T.ACCENT if step["now"] else T.RULE_C,
                  border_alpha=1.0 if step["now"] else card_style()["border_alpha"])
            if step["now"] and T.STYLE["card"] == "hairline":
                bar(slide, cx, y, 14 * S, 2 * S, T.ACCENT)
                bar(slide, cx, y, 2 * S, 14 * S, T.ACCENT)
        tx = cx + T.FLOW_PAD_X * S
        tw = cw - T.FLOW_PAD_X * 2 * S
        cy = y + T.FLOW_PAD_Y * S
        rich_text(slide, tx, cy, tw, t * 1.4, step["t"], t, S, font=T.MONO,
                  color=T.ACCENT, lh=1.4, spacing=0.12 * t)
        cy += t * 1.4 + 8 * S
        nh = rich_lines(step["n"], n, S, tw) * n * T.LH["fs_n"]
        rich_text(slide, tx, cy, tw, nh, step["n"], n, S, lh=T.LH["fs_n"])
    return h


# --------------------------------------------------------------------------
# cap-grid
# --------------------------------------------------------------------------


def _cap_cell(w, S):
    cols = T.CAP_COLS
    return (w - T.CAP_GAP * S * (cols - 1)) / cols


def m_cap_grid(b, w, S):
    cw = _cap_cell(w, S)
    icon = T.FS["cap_icon"] * S
    txt = T.FS["cap_text"] * S
    inner = cw - T.CAP_PAD_X * 2 * S
    rows = (len(b["caps"]) + T.CAP_COLS - 1) // T.CAP_COLS
    cell_h = 0.0
    for cap in b["caps"]:
        lines = T.wrap_lines(cap["text"], txt, inner)
        cell_h = max(cell_h, icon + 6 * S + lines * txt * T.LH["cap_text"]
                     + T.CAP_PAD_Y * 2 * S)
    return cell_h * rows + T.CAP_GAP * S * (rows - 1)


def d_cap_grid(slide, b, x, y, w, S):
    cw = _cap_cell(w, S)
    icon = T.FS["cap_icon"] * S
    txt = T.FS["cap_text"] * S
    inner = cw - T.CAP_PAD_X * 2 * S
    rows = (len(b["caps"]) + T.CAP_COLS - 1) // T.CAP_COLS
    cell_h = (m_cap_grid(b, w, S) - T.CAP_GAP * S * (rows - 1)) / rows

    for i, cap in enumerate(b["caps"]):
        r, c = divmod(i, T.CAP_COLS)
        cx = x + c * (cw + T.CAP_GAP * S)
        cy = y + r * (cell_h + T.CAP_GAP * S)
        card(slide, cx, cy, cw, cell_h, radius=T.RADIUS_SM * S,
             **card_style())
        ty = cy + T.CAP_PAD_Y * S
        plain_text(slide, cx, ty, cw, icon * 1.15, cap["icon"], icon * 0.8,
                   T.ACCENT, font=T.EMOJI, align="center", lh=1.0)
        ty += icon + 6 * S
        plain_text(slide, cx + T.CAP_PAD_X * S, ty, inner,
                   cell_h - (ty - cy), cap["text"], txt, T.MUTED, font=T.MONO,
                   align="center", lh=T.LH["cap_text"])
    return m_cap_grid(b, w, S)


# --------------------------------------------------------------------------
# ladder
# --------------------------------------------------------------------------


def _rung_h(rung, w, S):
    t = T.FS["rung_t"] * S
    inner = w - T.LADDER_PAD_X * 2 * S - (T.RUNG_DOT + T.RUNG_ICON_GAP) * S
    body = rich_lines(rung["runs"], t, S, inner) * t * T.LH["rung_t"]
    return max(body, T.RUNG_DOT * S) + T.LADDER_PAD_Y * 2 * S


def m_ladder(b, w, S):
    return (sum(_rung_h(r, w, S) for r in b["rungs"])
            + T.LADDER_GAP * S * max(0, len(b["rungs"]) - 1))


def d_ladder(slide, b, x, y, w, S):
    t = T.FS["rung_t"] * S
    nsz = T.FS["rung_n"] * S
    dot = T.RUNG_DOT * S
    cy = y
    for rung in b["rungs"]:
        h = _rung_h(rung, w, S)
        card(slide, x, cy, w, h, radius=T.RADIUS_SM * S,
             **card_style())
        dx = x + T.LADDER_PAD_X * S
        dy = cy + (h - dot) / 2
        if T.STYLE["rung_number"] == "block":
            bar(slide, x, cy, dot * 1.15, h, T.ACCENT)   # 줄 왼쪽을 꽉 채우는 블록
            dx = x + dot * 0.075                          # 블록 안에서 숫자를 가운데로
            dy = cy + (h - dot) / 2
            num_color = "FFFFFF"
        elif T.STYLE["rung_number"] == "filled":
            pill(slide, dx, dy, dot, dot, fill=T.INK)
            num_color = "FFFFFF"
        else:
            pill(slide, dx, dy, dot, dot, border=T.ACCENT)
            num_color = T.ACCENT
        plain_text(slide, dx, dy, dot, dot, rung["n"], nsz, num_color,
                   font=T.MONO, align="center", anchor="middle", bold=True, lh=1.0)
        tx = dx + dot + T.RUNG_ICON_GAP * S
        tw = w - (tx - x) - T.LADDER_PAD_X * S
        th = rich_lines(rung["runs"], t, S, tw) * t * T.LH["rung_t"]
        rich_text(slide, tx, cy + (h - th) / 2, tw, th, rung["runs"], t, S,
                  color=T.FG2, lh=T.LH["rung_t"])
        cy += h + T.LADDER_GAP * S
    return cy - y - T.LADDER_GAP * S


# --------------------------------------------------------------------------
# check-list
# --------------------------------------------------------------------------


def _chk_h(item, w, S):
    f = T.FS["chk_li"] * S
    inner = w - T.CHECK_PAD_X * 2 * S - (T.CHK_H + 14) * S
    body = rich_lines(item["runs"], f, S, inner) * f * T.LH["body"]
    return max(body, T.CHK_H * S) + T.CHECK_PAD_Y * 2 * S


def m_check_list(b, w, S):
    return (sum(_chk_h(i, w, S) for i in b["items"])
            + T.CHECK_GAP * S * max(0, len(b["items"]) - 1))


def d_check_list(slide, b, x, y, w, S):
    f = T.FS["chk_li"] * S
    nsz = T.FS["chk_n"] * S
    ch = T.CHK_H * S
    cy = y
    for item in b["items"]:
        h = _chk_h(item, w, S)
        card(slide, x, cy, w, h, radius=T.RADIUS_SM * S,
             **card_style())
        dx = x + T.CHECK_PAD_X * S
        dy = cy + (h - ch) / 2
        glyph_only = T.STYLE["chk"] == "accent2-glyph"
        if glyph_only:
            chip_bg, chip_fg = None, T.ACCENT2
        elif T.STYLE["chk"] == "accent2-fill":
            chip_bg, chip_fg = T.ACCENT2, "FFFFFF"
        else:
            chip_bg, chip_fg = T.ACCENT, T.INK
        if chip_bg:
            pill(slide, dx, dy, ch * 1.15, ch, fill=chip_bg)
        plain_text(slide, dx, dy, ch * 1.15, ch, item["n"], nsz, chip_fg,
                   font=T.MONO, align="center", anchor="middle", bold=True, lh=1.0)
        tx = dx + ch * 1.15 + 14 * S
        tw = w - (tx - x) - T.CHECK_PAD_X * S
        th = rich_lines(item["runs"], f, S, tw) * f * T.LH["body"]
        rich_text(slide, tx, cy + (h - th) / 2, tw, th, item["runs"], f, S,
                  color=T.FG2, lh=T.LH["body"])
        cy += h + T.CHECK_GAP * S
    return cy - y - T.CHECK_GAP * S


# --------------------------------------------------------------------------
# stat-row
# --------------------------------------------------------------------------


def m_stat_row(b, w, S):
    num = T.FS["stat_num"] * S
    lab = T.FS["stat_lab"] * S
    k = len(b["stats"])
    cw = (w - T.STAT_GAP * S * (k - 1)) / k
    inner = cw - T.STAT_PAD_X * 2 * S
    lab_lines = max(T.wrap_lines(s["lab"], lab, inner) for s in b["stats"])
    return num + 8 * S + lab_lines * lab * T.LH["stat_lab"] + T.STAT_PAD_Y * 2 * S


def d_stat_row(slide, b, x, y, w, S):
    num = T.FS["stat_num"] * S
    lab = T.FS["stat_lab"] * S
    k = len(b["stats"])
    cw = (w - T.STAT_GAP * S * (k - 1)) / k
    h = m_stat_row(b, w, S)
    for i, st in enumerate(b["stats"]):
        cx = x + i * (cw + T.STAT_GAP * S)
        card(slide, cx, y, cw, h, radius=T.RADIUS_SM * S,
             **card_style())
        num_col = T.INK if T.STYLE["stat_num"] == "ink" else T.ACCENT
        plain_text(slide, cx, y + T.STAT_PAD_Y * S, cw, num * 1.05, st["num"],
                   num, num_col, font=T.SANS_BLACK, align="center", bold=True, lh=1.0)
        plain_text(slide, cx + T.STAT_PAD_X * S, y + T.STAT_PAD_Y * S + num + 8 * S,
                   cw - T.STAT_PAD_X * 2 * S, h, st["lab"], lab, T.MUTED,
                   font=T.MONO, align="center", lh=T.LH["stat_lab"])
    return h


# --------------------------------------------------------------------------
# case-card
# --------------------------------------------------------------------------


def m_case_card(b, w, S):
    inner = w - T.CASE_PAD_X * 2 * S
    nm = T.FS["case_name"] * S
    h = nm * 1.4
    if b["what"]:
        wt = T.FS["case_what"] * S
        h += 8 * S + rich_lines(b["what"], wt, S, inner) * wt * T.LH["case_what"]
    if b["lesson"]:
        ls = T.FS["case_lesson"] * S
        h += 6 * S + rich_lines(b["lesson"], ls, S, inner) * ls * 1.4
    return h + T.CASE_PAD_Y * 2 * S


def d_case_card(slide, b, x, y, w, S):
    h = m_case_card(b, w, S)
    tcard(slide, x, y, w, h, radius=T.RADIUS_SM * S)
    if T.STYLE["case_card"] == "left-bar":
        bar(slide, x, y, 5 * S, h, T.ACCENT)
    else:
        corner_brackets(slide, x, y, w, h, T.ACCENT, arm=T.BRACKET * S, thick=2 * S)

    tx = x + T.CASE_PAD_X * S
    tw = w - T.CASE_PAD_X * 2 * S
    cy = y + T.CASE_PAD_Y * S
    nm = T.FS["case_name"] * S
    rich_text(slide, tx, cy, tw * 0.68, nm * 1.4, b["name"], nm, S,
              color=T.FG_STRONG, lh=1.4, bold=True)
    if b["rev"]:
        rv = T.FS["case_rev"] * S
        if T.STYLE["case_rev"] == "pill":
            pw = T.text_width_px(b["rev"], rv) + 24 * S
            ph = rv * 1.6
            pill(slide, x + w - T.CASE_PAD_X * S - pw, cy, pw, ph,
                 fill=T.ACCENT2, fill_alpha=T.A_ACCENT_06)
            plain_text(slide, x + w - T.CASE_PAD_X * S - pw, cy, pw, ph,
                       b["rev"], rv, T.ACCENT2, font=T.MONO, align="center",
                       anchor="middle", bold=True, lh=1.0)
        else:
            plain_text(slide, tx + tw * 0.68, cy, tw * 0.32, rv * 1.4, b["rev"], rv,
                       T.ACCENT, font=T.MONO, align="right", bold=True, lh=1.4)
    cy += nm * 1.4
    if b["what"]:
        wt = T.FS["case_what"] * S
        wh = rich_lines(b["what"], wt, S, tw) * wt * T.LH["case_what"]
        rich_text(slide, tx, cy + 8 * S, tw, wh, b["what"], wt, S,
                  color=T.FG2, lh=T.LH["case_what"])
        cy += 8 * S + wh
    if b["lesson"]:
        ls = T.FS["case_lesson"] * S
        lh_ = rich_lines(b["lesson"], ls, S, tw) * ls * 1.4
        if T.STYLE["case_lesson"] == "bold":
            rich_text(slide, tx, cy + 6 * S, tw, lh_, b["lesson"], ls, S,
                      font=T.SANS, color=T.ACCENT, lh=1.4, bold=True)
        else:
            rich_text(slide, tx, cy + 6 * S, tw, lh_, b["lesson"], ls, S,
                      font=T.SERIF, color=T.ACCENT, lh=1.4)
    return h


# --------------------------------------------------------------------------
# brief-block / pre>code — 모노 블록
# --------------------------------------------------------------------------


def _mono_h(text, w, S, pad_x, pad_y):
    f = T.FS["code"] * S
    lines = T.wrap_lines(text.strip("\n"), f, w - pad_x * 2 * S)
    return lines * f * T.LH["code"] + pad_y * 2 * S


def m_brief(b, w, S):
    return _mono_h(b["text"], w, S, T.BRIEF_PAD_X, T.BRIEF_PAD_Y)


def m_code(b, w, S):
    return _mono_h(b["text"], w, S, T.BRIEF_PAD_X, T.BRIEF_PAD_Y)


def _draw_mono(slide, text, x, y, w, S, accent_edge):
    h = _mono_h(text, w, S, T.BRIEF_PAD_X, T.BRIEF_PAD_Y)
    if T.CODE_BG:
        # 밝은 테마: 코드·지시문은 어두운 판으로 뒤집어 대비를 만든다
        tcard(slide, x, y, w, h, radius=T.RADIUS_SM * S,
              fill=T.CODE_BG, fill_alpha=1.0, border=T.CODE_BG, border_alpha=0.0)
    else:
        tcard(slide, x, y, w, h, radius=T.RADIUS_SM * S)
        if accent_edge:
            bar(slide, x, y, 2 * S, h, T.ACCENT)
    f = T.FS["code"] * S
    tb = textbox(slide, x + T.BRIEF_PAD_X * S, y + T.BRIEF_PAD_Y * S,
                 w - T.BRIEF_PAD_X * 2 * S, h - T.BRIEF_PAD_Y * 2 * S)
    first = True
    for line in text.strip("\n").split("\n"):
        p = para(tb.text_frame, first=first, line_spacing=ls(f, T.LH["code"]))
        run(p, line or " ", T.MONO, f, T.CODE_FG)
        first = False
    return h


def d_brief(slide, b, x, y, w, S):
    return _draw_mono(slide, b["text"], x, y, w, S, accent_edge=True)


def d_code(slide, b, x, y, w, S):
    return _draw_mono(slide, b["text"], x, y, w, S, accent_edge=True)


# --------------------------------------------------------------------------
# source-list
# --------------------------------------------------------------------------


def _src_h(item, w, S):
    f = T.FS["src"] * S
    inner = w - T.SRC_PAD_X * 2 * S
    head = f"{item['id']} {item['title']}"
    tail = f"{item['meta']} [원문]"
    lines = (T.wrap_lines(head, f, inner) + T.wrap_lines(tail, f, inner))
    return lines * f * T.LH["src"] + T.SRC_PAD_Y * 2 * S


def m_source_list(b, w, S):
    return (sum(_src_h(i, w, S) for i in b["items"])
            + T.SRC_GAP * S * max(0, len(b["items"]) - 1))


def d_source_list(slide, b, x, y, w, S):
    f = T.FS["src"] * S
    cy = y
    for item in b["items"]:
        h = _src_h(item, w, S)
        shadow = T.STYLE["card"] == "shadow"
        # 출처 카드는 헤어라인 테마에서 테두리가 없다 (좌측 바만으로 구분한다)
        tcard(slide, x, cy, w, h,
              radius=(T.RADIUS_SM * S if shadow else 0),
              border=T.RULE_C if shadow else None, border_alpha=0.0)
        bar(slide, x, cy, (4 if shadow else 2) * S, h, T.ACCENT)
        tb = textbox(slide, x + T.SRC_PAD_X * S, cy + T.SRC_PAD_Y * S,
                     w - T.SRC_PAD_X * 2 * S, h - T.SRC_PAD_Y * 2 * S)
        p = para(tb.text_frame, first=True, line_spacing=ls(f, T.LH["src"]))
        run(p, item["id"] + " ", T.MONO, f, T.ACCENT, bold=True)
        run(p, item["title"], T.SANS, f, T.FG_STRONG, bold=True)
        line_break(p)
        run(p, item["meta"] + "  ", T.SANS, f, T.MUTED)
        if item["url"]:
            run(p, "[원문]", T.SANS, f, T.ACCENT, link=item["url"])
        cy += h + T.SRC_GAP * S
    return cy - y - T.SRC_GAP * S


# --------------------------------------------------------------------------
# title-block
# --------------------------------------------------------------------------


def m_title_block(b, w, S):
    h = 0.0
    if b["eyebrow"]:
        e = T.FS["eyebrow_title"] * S
        h += e * 1.4 + 0.7 * e
    h1 = T.FS["h1_title"] * S
    h += rich_lines(b["h1"], h1, S, w, T.SANS) * h1 * T.LH["h1"]
    if b["subtitle"]:
        s = T.FS["subtitle"] * S
        h += 0.4 * s + rich_lines(b["subtitle"], s, S, w) * s * T.LH["subtitle"]
    return h


def d_title_block(slide, b, x, y, w, S):
    cy = y
    if b["eyebrow"]:
        e = T.FS["eyebrow_title"] * S
        runs = b["eyebrow"]
        if T.STYLE["eyebrow_prefix"]:
            runs = [{"text": T.STYLE["eyebrow_prefix"], "style": PLAIN,
                     "link": None}] + runs
        rich_text(slide, x, cy, w, e * 1.4, runs, e, S, font=T.MONO,
                  color=T.ACCENT, lh=1.4, spacing=0.2 * e)
        cy += e * 1.4 + 0.7 * e
    h1 = T.FS["h1_title"] * S
    hh = rich_lines(b["h1"], h1, S, w, T.SANS) * h1 * T.LH["h1"]
    rich_text(slide, x, cy, w, hh, b["h1"], h1, S, font=T.SANS_BLACK,
              color=T.FG_STRONG, lh=T.LH["h1"], bold=True)
    cy += hh
    if b["subtitle"]:
        s = T.FS["subtitle"] * S
        sh = rich_lines(b["subtitle"], s, S, w) * s * T.LH["subtitle"]
        rich_text(slide, x, cy + 0.4 * s, w, sh, b["subtitle"], s, S,
                  color=T.MUTED, lh=T.LH["subtitle"])
    return m_title_block(b, w, S)


# --------------------------------------------------------------------------
# card-grid — grid-3 / v2-card 류 (제목 + 한 줄 카드 격자)
# --------------------------------------------------------------------------


def _grid_cols(n):
    return min(n, 4) if n else 1


def _card_grid_geom(b, w, S):
    cols = _grid_cols(len(b["cards"]))
    cw = (w - T.GRID_GAP * S * (cols - 1)) / cols
    rows = (len(b["cards"]) + cols - 1) // cols
    ti = T.FS["card_title"] * S
    bo = T.FS["card_body"] * S
    inner = cw - T.CARD_PAD_X * 2 * S
    cell = 0.0
    for c in b["cards"]:
        h = (rich_lines(c["title"], ti, S, inner) * ti * 1.4
             + 6 * S
             + rich_lines(c["body"], bo, S, inner) * bo * T.LH["card_body"]
             + T.CARD_PAD_Y * 2 * S)
        cell = max(cell, h)
    return cols, cw, rows, cell


def m_card_grid(b, w, S):
    _, _, rows, cell = _card_grid_geom(b, w, S)
    return cell * rows + T.GRID_GAP * S * (rows - 1)


def d_card_grid(slide, b, x, y, w, S):
    cols, cw, rows, cell = _card_grid_geom(b, w, S)
    ti = T.FS["card_title"] * S
    bo = T.FS["card_body"] * S
    inner = cw - T.CARD_PAD_X * 2 * S
    for i, c in enumerate(b["cards"]):
        r, col = divmod(i, cols)
        cx = x + col * (cw + T.GRID_GAP * S)
        cy = y + r * (cell + T.GRID_GAP * S)
        card(slide, cx, cy, cw, cell, radius=T.RADIUS_SM * S,
             **card_style())
        ty = cy + T.CARD_PAD_Y * S
        th = rich_lines(c["title"], ti, S, inner) * ti * 1.4
        rich_text(slide, cx + T.CARD_PAD_X * S, ty, inner, th, c["title"], ti, S,
                  color=T.FG_STRONG, lh=1.4, bold=True)
        bh = rich_lines(c["body"], bo, S, inner) * bo * T.LH["card_body"]
        rich_text(slide, cx + T.CARD_PAD_X * S, ty + th + 6 * S, inner, bh,
                  c["body"], bo, S, color=T.FG2, lh=T.LH["card_body"])
    return m_card_grid(b, w, S)


# --------------------------------------------------------------------------
# table — 편집 가능한 네이티브 표 (trouble-table 포함)
# --------------------------------------------------------------------------


def _table_geom(b, w, S):
    ncols = max(len(r) for r in b["rows"])
    cw = w / ncols
    f = T.FS["cell"] * S
    inner = cw - T.CELL_PAD_X * 2 * S
    heights = []
    for row in b["rows"]:
        lines = max(rich_lines(c["runs"], f, S, inner) for c in row)
        heights.append(lines * f * T.LH["cell"] + T.CELL_PAD_Y * 2 * S)
    return ncols, cw, heights


def m_table(b, w, S):
    _, _, heights = _table_geom(b, w, S)
    return sum(heights)


def d_table(slide, b, x, y, w, S):
    ncols, cw, heights = _table_geom(b, w, S)
    f = T.FS["cell"] * S
    gf = slide.shapes.add_table(len(b["rows"]), ncols, px(x), px(y), px(w),
                                px(sum(heights)))
    tbl = gf.table
    style_table_dark(tbl, T.RULE_C, T.A_RULE)
    for i in range(ncols):
        tbl.columns[i].width = px(cw)
    for ri, row in enumerate(b["rows"]):
        tbl.rows[ri].height = px(heights[ri])
        for ci in range(ncols):
            cell = tbl.cell(ri, ci)
            src = row[ci] if ci < len(row) else None
            cell.margin_left = px(T.CELL_PAD_X * S)
            cell.margin_right = px(T.CELL_PAD_X * S)
            cell.margin_top = px(T.CELL_PAD_Y * S)
            cell.margin_bottom = px(T.CELL_PAD_Y * S)
            fill_rgba(cell, T.SURFACE, T.A_SURFACE_1)
            tf = cell.text_frame
            tf.word_wrap = True
            p = para(tf, first=True, line_spacing=ls(f, T.LH["cell"]))
            if src:
                color = T.MUTED if src["muted"] else T.FG2
                for r in src["runs"]:
                    if r["style"] == BR:
                        line_break(p)
                        continue
                    fo, size, col, bo, it = run_style(
                        r["style"], f, S, T.SANS, color, r["text"],
                        ctx_bold=src["th"])
                    run(p, r["text"], fo, size, col, bold=bo, italic=it,
                        link=r.get("link"))
    return sum(heights)


# --------------------------------------------------------------------------
# tier-grid — 요금제
# --------------------------------------------------------------------------


def _tier_geom(b, w, S):
    k = len(b["tiers"])
    cw = (w - T.GRID_GAP * S * (k - 1)) / k
    nm = T.FS["tier_name"] * S
    pr = T.FS["tier_price"] * S
    li = T.FS["tier_li"] * S
    inner = cw - T.TIER_PAD_X * 2 * S
    h = 0.0
    for t in b["tiers"]:
        th = nm * 1.4 + 6 * S + rich_lines(t["price"], pr, S, inner) * pr * 1.2 + 8 * S
        for it in t["items"]:
            th += rich_lines(it, li, S, inner) * li * 1.5 + 5 * S
        h = max(h, th + T.TIER_PAD_Y * 2 * S)
    return cw, h


def m_tier_grid(b, w, S):
    return _tier_geom(b, w, S)[1]


def d_tier_grid(slide, b, x, y, w, S):
    cw, h = _tier_geom(b, w, S)
    nm = T.FS["tier_name"] * S
    pr = T.FS["tier_price"] * S
    li = T.FS["tier_li"] * S
    inner = cw - T.TIER_PAD_X * 2 * S
    for i, t in enumerate(b["tiers"]):
        cx = x + i * (cw + T.GRID_GAP * S)
        tcard(slide, cx, y, cw, h, radius=T.RADIUS_SM * S,
              border=T.ACCENT if t["featured"] else T.RULE_C,
              border_alpha=1.0 if t["featured"] else card_style()["border_alpha"])
        tx = cx + T.TIER_PAD_X * S
        cy = y + T.TIER_PAD_Y * S
        plain_text(slide, tx, cy, inner, nm * 1.4, t["name"].upper(), nm, T.ACCENT,
                   font=T.MONO, lh=1.4, spacing=0.12 * nm)
        cy += nm * 1.4 + 6 * S
        ph = rich_lines(t["price"], pr, S, inner) * pr * 1.2
        rich_text(slide, tx, cy, inner, ph, t["price"], pr, S,
                  color=T.FG_STRONG, lh=1.2, bold=True)
        cy += ph + 8 * S
        for it in t["items"]:
            ih = rich_lines(it, li, S, inner) * li * 1.5
            rich_text(slide, tx, cy, inner, ih, it, li, S, color=T.MUTED, lh=1.5)
            cy += ih + 5 * S
    return h


# --------------------------------------------------------------------------
# chip-row — 알약 태그
# --------------------------------------------------------------------------


def _chip_layout(b, w, S):
    """(x, y, width, chip) 목록과 전체 높이."""
    f = T.FS["chip"] * S
    ch = f * 1.5 + T.CHIP_PAD_Y * 2 * S
    out, cx, cy = [], 0.0, 0.0
    for c in b["chips"]:
        cwid = (T.text_width_px(c["text"], f)
                + (0 if c["arrow"] else T.CHIP_PAD_X * 2 * S))
        if cx > 0 and cx + cwid > w:
            cx, cy = 0.0, cy + ch + T.CHIP_GAP * S
        out.append((cx, cy, cwid, c))
        cx += cwid + T.CHIP_GAP * S
    return out, cy + ch, ch


def m_chip_row(b, w, S):
    return _chip_layout(b, w, S)[1]


def d_chip_row(slide, b, x, y, w, S):
    items, total, ch = _chip_layout(b, w, S)
    f = T.FS["chip"] * S
    for cx, cy, cwid, c in items:
        if c["arrow"]:
            plain_text(slide, x + cx, y + cy, cwid, ch, c["text"], f * 1.2,
                       T.MUTED, align="center", anchor="middle", lh=1.0)
            continue
        pill(slide, x + cx, y + cy, cwid, ch, fill=T.ACCENT,
             fill_alpha=T.A_ACCENT_06, border=T.ACCENT, border_alpha=T.A_ACCENT_25)
        plain_text(slide, x + cx, y + cy, cwid, ch, c["text"], f, T.ACCENT,
                   font=T.MONO, align="center", anchor="middle", lh=1.0)
    return total


# --------------------------------------------------------------------------
# qr-grid — 마무리 QR 카드
# --------------------------------------------------------------------------


def _qr_geom(b, w, S):
    cw = (w - T.QR_GAP * S) / 2 if len(b["cards"]) > 1 else w
    eb = T.FS["qr_eyebrow"] * S
    nm = T.FS["qr_name"] * S
    de = T.FS["qr_desc"] * S
    ur = T.FS["qr_url"] * S
    body_w = cw - T.QR_PAD_X * 2 * S - T.QR_SIZE * S - T.QR_INNER_GAP * S
    h = 0.0
    for c in b["cards"]:
        bh = (eb * 1.4 + 5 * S
              + rich_lines(c["name"], nm, S, body_w) * nm * 1.15 + 8 * S
              + rich_lines(c["desc"], de, S, body_w) * de * T.LH["qr_desc"] + 9 * S
              + T.wrap_lines(c["url_text"], ur, body_w) * ur * 1.4)
        h = max(h, max(bh, T.QR_SIZE * S) + T.QR_PAD_Y * 2 * S)
    return cw, body_w, h


def m_qr_grid(b, w, S, base_dir=None):
    return _qr_geom(b, w, S)[2]


def d_qr_grid(slide, b, x, y, w, S, base_dir=None):
    cw, body_w, h = _qr_geom(b, w, S)
    eb = T.FS["qr_eyebrow"] * S
    nm = T.FS["qr_name"] * S
    de = T.FS["qr_desc"] * S
    ur = T.FS["qr_url"] * S
    for i, c in enumerate(b["cards"]):
        cx = x + i * (cw + T.QR_GAP * S)
        card(slide, cx, y, cw, h, radius=14 * S,
             **card_style())
        qx = cx + T.QR_PAD_X * S
        qy = y + (h - T.QR_SIZE * S) / 2
        card(slide, qx, qy, T.QR_SIZE * S, T.QR_SIZE * S, radius=12 * S,
             fill="FFFFFF")
        img = resolve_asset(c["src"], base_dir)
        if img:
            pad = 9 * S
            slide.shapes.add_picture(str(img), px(qx + pad), px(qy + pad),
                                     px(T.QR_SIZE * S - pad * 2),
                                     px(T.QR_SIZE * S - pad * 2))
        tx = qx + T.QR_SIZE * S + T.QR_INNER_GAP * S
        cy = y + T.QR_PAD_Y * S
        plain_text(slide, tx, cy, body_w, eb * 1.4, c["eyebrow"].upper(), eb,
                   T.ACCENT, font=T.MONO, lh=1.4, spacing=0.12 * eb)
        cy += eb * 1.4 + 5 * S
        nh = rich_lines(c["name"], nm, S, body_w) * nm * 1.15
        rich_text(slide, tx, cy, body_w, nh, c["name"], nm, S,
                  color=T.FG_STRONG, lh=1.15, bold=True)
        cy += nh + 8 * S
        dh = rich_lines(c["desc"], de, S, body_w) * de * T.LH["qr_desc"]
        rich_text(slide, tx, cy, body_w, dh, c["desc"], de, S, color=T.FG2,
                  lh=T.LH["qr_desc"])
        cy += dh + 9 * S
        plain_text(slide, tx, cy, body_w, ur * 1.4 * 3, c["url_text"], ur,
                   T.ACCENT, font=T.MONO, lh=1.4)
    return h


# --------------------------------------------------------------------------
# image — image-slot / 본문 <img>
# --------------------------------------------------------------------------

MAX_IMG_H = 420          # 본문 이미지가 슬라이드를 다 먹지 않도록


def resolve_asset(src, base_dir):
    """상대 경로 이미지를 덱 폴더 기준으로 찾는다. 없으면 None."""
    if not src or src.startswith(("http://", "https://", "data:")):
        return None
    if base_dir is None:
        return None
    p = (Path(base_dir) / src).resolve()
    return p if p.exists() else None


def _img_size(b, w, S, base_dir):
    p = resolve_asset(b["src"], base_dir)
    if p is None:
        return None, w, T.FS["small"] * S * 2
    with Image.open(p) as im:
        iw, ih = im.size
    scale = min(w / iw, MAX_IMG_H * S / ih, 1.0)
    return p, iw * scale, ih * scale


def m_image(b, w, S, base_dir=None):
    return _img_size(b, w, S, base_dir)[2]


def d_image(slide, b, x, y, w, S, base_dir=None):
    p, iw, ih = _img_size(b, w, S, base_dir)
    if p is None:
        # 파일을 못 찾으면 자리와 alt 를 남긴다 (조용히 사라지지 않게)
        card(slide, x, y, w, ih, radius=T.RADIUS_SM * S,
             **card_style())
        plain_text(slide, x, y, w, ih, f"[이미지 없음] {b['src']}",
                   T.FS["small"] * S, T.MUTED, font=T.MONO, align="center",
                   anchor="middle", lh=1.4)
        return ih
    slide.shapes.add_picture(str(p), px(x + (w - iw) / 2), px(y), px(iw), px(ih))
    return ih


# --------------------------------------------------------------------------
# 디스패치
# --------------------------------------------------------------------------

MEASURE = {
    "h1": m_h, "h2": m_h, "p": m_p, "bullets": m_bullets, "tl": m_tl,
    "dual": m_dual, "card-row": m_card_row, "thesis-flow": m_thesis_flow,
    "cap-grid": m_cap_grid, "ladder": m_ladder, "check-list": m_check_list,
    "stat-row": m_stat_row, "case-card": m_case_card, "brief": m_brief,
    "code": m_code, "source-list": m_source_list, "title-block": m_title_block,
    "card-grid": m_card_grid, "table": m_table, "tier-grid": m_tier_grid,
    "chip-row": m_chip_row, "qr-grid": m_qr_grid, "image": m_image,
}

DRAW = {
    "h1": d_h, "h2": d_h, "p": d_p, "bullets": d_bullets, "tl": d_tl,
    "dual": d_dual, "card-row": d_card_row, "thesis-flow": d_thesis_flow,
    "cap-grid": d_cap_grid, "ladder": d_ladder, "check-list": d_check_list,
    "stat-row": d_stat_row, "case-card": d_case_card, "brief": d_brief,
    "code": d_code, "source-list": d_source_list, "title-block": d_title_block,
    "card-grid": d_card_grid, "table": d_table, "tier-grid": d_tier_grid,
    "chip-row": d_chip_row, "qr-grid": d_qr_grid, "image": d_image,
}

# 덱 폴더 기준으로 이미지를 찾아야 하는 블록
NEEDS_BASE_DIR = {"image", "qr-grid"}


def measure(b, w, S, base_dir=None):
    fn = MEASURE[b["kind"]]
    return fn(b, w, S, base_dir) if b["kind"] in NEEDS_BASE_DIR else fn(b, w, S)


def draw(slide, b, x, y, w, S, base_dir=None):
    fn = DRAW[b["kind"]]
    if b["kind"] in NEEDS_BASE_DIR:
        return fn(slide, b, x, y, w, S, base_dir)
    return fn(slide, b, x, y, w, S)
