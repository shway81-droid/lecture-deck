"""lecture-deck 의 index.html → 텍스트가 편집되는 PPTX.

    python build.py <index.html> [out.pptx] [--notes] [--quiet]

--notes 를 주면 `<aside class="notes">` 를 파워포인트 발표자 노트로 넣는다
(강사 교안용). 기본은 노트 없는 배포용이다.
"""

import argparse
import sys
from pathlib import Path

import preflight

preflight.require()

from pptx import Presentation  # noqa: E402

import background  # noqa: E402
import themes  # noqa: E402
import render  # noqa: E402
import theme as T  # noqa: E402
from parse import parse_deck  # noqa: E402
from shapes import bar, px, set_theme_hyperlink_color  # noqa: E402

# reveal 의 shrink-to-fit 을 흉내낸 축소 단계
SCALES = [1.0 - i * 0.02 for i in range(15)]   # 1.00 → 0.72


def _is_footer_meta(b) -> bool:
    return b["kind"] == "p" and "footer-meta" in b["classes"]


def _has_class(b, name) -> bool:
    return b["kind"] == "p" and name in b["classes"]


def _stack_height(blocks, w, S, base_dir) -> float:
    """블록을 쌓았을 때 총 높이. 인접 margin 은 CSS 처럼 겹친다."""
    total = 0.0
    prev_mb = None
    for b in blocks:
        mt, mb = render.margins(b, S)
        if prev_mb is not None:
            total += max(prev_mb, mt)
        total += render.measure(b, w, S, base_dir)
        prev_mb = mb
    return total


def _layout(blocks, w, base_dir):
    for S in SCALES:
        h = _stack_height(blocks, w, S, base_dir)
        if h <= T.CONTENT_H:
            return S, h
    S = SCALES[-1]
    return S, _stack_height(blocks, w, S, base_dir)


def _draw_stack(slide, blocks, x, y, w, S, base_dir):
    prev_mb = None
    for b in blocks:
        mt, mb = render.margins(b, S)
        if prev_mb is not None:
            y += max(prev_mb, mt)
        render.draw(slide, b, x, y, w, S, base_dir)
        y += render.measure(b, w, S, base_dir)
        prev_mb = mb
    return y


def _slide_background(slide, s, base_dir, bg_png):
    """data-background-image 가 있으면 그걸, 없으면 덱 기본 배경을 깐다."""
    img = render.resolve_asset(s["background"], base_dir) if s["background"] else None
    if img is None:
        slide.shapes.add_picture(str(bg_png), 0, 0, px(T.CANVAS_W), px(T.CANVAS_H))
        return
    slide.shapes.add_picture(str(img), 0, 0, px(T.CANVAS_W), px(T.CANVAS_H))
    # data-background-opacity 는 검정 오버레이로 흉내낸다
    try:
        op = float(s["background_opacity"]) if s["background_opacity"] else 1.0
    except ValueError:
        op = 1.0
    if op < 1.0:
        bar(slide, 0, 0, T.CANVAS_W, T.CANVAS_H, T.BG, 1.0 - op)


def _render_hero_img(slide, s, base_dir, S=1.0):
    """전면 이미지 + 왼쪽 아래 한 줄 (theme.css .hero-img)."""
    subs = [b for b in s["blocks"] if _has_class(b, "hero-sub")]
    metas = [b for b in s["blocks"] if _is_footer_meta(b)]
    rest = [b for b in s["blocks"] if b not in subs and b not in metas]

    w = T.CANVAS_W * 0.46
    y = T.CANVAS_H - T.HERO_SUB_BOTTOM
    for b in reversed(subs):
        h = render.measure(b, w, S, base_dir)
        y -= h
        render.draw(slide, b, T.HERO_SUB_LEFT, y, w, S, base_dir)
    for b in metas:
        h = render.measure(b, T.CONTENT_W, S, base_dir)
        render.draw(slide, b, T.HERO_SUB_LEFT, T.CANVAS_H - 40 - h,
                    T.CONTENT_W, S, base_dir)
    if rest:
        S2, _ = _layout(rest, T.CONTENT_W, base_dir)
        _draw_stack(slide, rest, T.PAD_X, T.PAD_Y, T.CONTENT_W, S2, base_dir)


def _render_concept(slide, s, base_dir, scrim_png):
    """왼쪽 글 + 오른쪽으로 사라지는 스크림 (theme.css .concept)."""
    slide.shapes.add_picture(str(scrim_png), 0, 0, px(T.CANVAS_W), px(T.CANVAS_H))

    w = T.CANVAS_W * 0.50 - T.BEAT_LEFT
    S, total = _layout(s["blocks"], w, base_dir)
    _draw_stack(slide, s["blocks"], T.BEAT_LEFT, (T.CANVAS_H - total) / 2, w, S,
                base_dir)


def build(html: Path, out: Path, keep_notes=False, theme_name=None) -> dict:
    html = Path(html)
    base_dir = html.parent

    # 테마: 명시값 → 덱의 theme.css 자동 감지 → 기본값.
    # 감지에 실패하면 조용히 넘어가지 않고 호출부가 알릴 수 있게 표시를 돌려준다.
    active, detected = T.activate(theme_name, base_dir / "theme.css")

    deck = parse_deck(html, keep_notes=keep_notes)

    # 배경은 PPTX 안에 embed 되므로 이 파일들은 재생성을 아끼기 위한 캐시일 뿐이다.
    # 테마마다 다른 그림이라 파일명을 테마로 갈라둔다.
    cache = Path(__file__).parent / "cache"
    bg_png = cache / f"pptx-bg-{T.THEME_NAME}.png"
    scrim_png = cache / f"pptx-scrim-{T.SCRIM}.png"
    if not bg_png.exists():
        {"warm": background.build_warm, "slate": background.build_slate,
         "grid": background.build_grid, "dot": background.build_dot}.get(T.BACKGROUND, background.build)(bg_png)
    if not scrim_png.exists():
        background.build_scrim(scrim_png)

    prs = Presentation()
    prs.slide_width = px(T.CANVAS_W)
    prs.slide_height = px(T.CANVAS_H)
    set_theme_hyperlink_color(prs, T.ACCENT)
    blank = prs.slide_layouts[6]

    stats = {"slides": 0, "shrunk": [], "overflow": [], "unknown": {},
             "missing_images": [], "theme": active,
             "theme_detected": detected or bool(theme_name)}

    for s in deck:
        slide = prs.slides.add_slide(blank)
        _slide_background(slide, s, base_dir, bg_png)

        for u in s["unknown"]:
            stats["unknown"].setdefault(u, []).append(s["index"])
        for b in s["blocks"]:
            if b["kind"] == "image" and render.resolve_asset(b["src"], base_dir) is None:
                stats["missing_images"].append((s["index"], b["src"]))

        if s["kind"] == "hero-img":
            _render_hero_img(slide, s, base_dir)
        elif s["kind"] == "concept":
            for b in s["blocks"]:
                if b["kind"] == "h2" and not b.get("variant"):
                    b["variant"] = "beat"
            _render_concept(slide, s, base_dir, scrim_png)
        else:
            blocks = list(s["blocks"])
            if s["kind"] == "divider-slide":
                for b in blocks:
                    if b["kind"] == "h1":
                        b["variant"] = "divider"

            pinned = None
            if s["kind"] in ("title-slide", "divider-slide") and blocks \
                    and _is_footer_meta(blocks[-1]):
                pinned = blocks.pop()

            S, total = _layout(blocks, T.CONTENT_W, base_dir)
            if S < 1.0:
                stats["shrunk"].append((s["index"], round(S, 2)))
            if total > T.CONTENT_H:
                stats["overflow"].append((s["index"], round(total - T.CONTENT_H)))

            y = T.PAD_Y + max(0.0, (T.CONTENT_H - total) / 2)
            _draw_stack(slide, blocks, T.PAD_X, y, T.CONTENT_W, S, base_dir)

            if pinned is not None:
                fh = render.measure(pinned, T.CONTENT_W, S, base_dir)
                render.draw(slide, pinned, T.PAD_X, T.CANVAS_H - T.PAD_Y - fh,
                            T.CONTENT_W, S, base_dir)

        if keep_notes and s["notes"]:
            slide.notes_slide.notes_text_frame.text = s["notes"]

        stats["slides"] += 1

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return stats


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="lecture-deck index.html → 편집 가능 PPTX")
    ap.add_argument("html", type=Path, nargs="?", help="덱의 index.html")
    ap.add_argument("out", type=Path, nargs="?", help="출력 .pptx (기본: 덱 폴더 이름)")
    ap.add_argument("--notes", action="store_true",
                    help="발표자 노트를 PPTX 노트에 넣는다 (강사 교안용)")
    ap.add_argument("--quiet", action="store_true")
    ap.add_argument("--check", action="store_true",
                    help="의존성과 폰트만 확인하고 끝낸다")
    ap.add_argument("--theme", default=None,
                    help="테마 강제 지정. 생략하면 덱의 theme.css 를 보고 고른다 "
                         f"({', '.join(sorted(themes.REGISTRY))})")
    a = ap.parse_args(argv)

    if a.check:
        print(f"python: {sys.version.split()[0]}  ({sys.executable})")
        return preflight.report()
    if a.html is None:
        ap.error("index.html 경로가 필요합니다 (또는 --check)")

    out = a.out or a.html.parent / f"{a.html.parent.name}.pptx"
    st = build(a.html, out, keep_notes=a.notes, theme_name=a.theme)

    print(f"슬라이드 {st['slides']}장 → {out}")
    if not a.quiet:
        src = "지정" if a.theme else ("theme.css 감지" if st["theme_detected"]
                                      else "기본값, theme.css 를 못 읽음")
        print(f"테마: {st['theme']}  ({src})")
        print(f"폰트: {T.font_report()}")
    if st["shrunk"]:
        print("  축소 적용: " + ", ".join(f"S{i}={v}" for i, v in st["shrunk"]))
    if st["overflow"]:
        print("  !! 여전히 넘침: " + ", ".join(f"S{i}=+{v}px" for i, v in st["overflow"]))
    if st["missing_images"]:
        print("  !! 이미지 파일 없음: "
              + ", ".join(f"S{i}:{p}" for i, p in st["missing_images"]))
    if st["unknown"]:
        print("  전용 렌더러가 없어 글자만 살린 요소 (필요하면 렌더러 추가):")
        for sel, idxs in st["unknown"].items():
            print(f"    {sel} — S{', S'.join(map(str, idxs))}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
